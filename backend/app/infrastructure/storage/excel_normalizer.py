"""Read-only нормализация Excel-выгрузок IEK и Systeme Electric."""

from __future__ import annotations

import re
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.domain.entities import BulkOrderEvent, Inbound, MonthPoint, Sku
from app.domain.exceptions import StorageError
from app.infrastructure.storage.data_quality import DataQualityReport
from app.infrastructure.storage.memory_repo import MemorySkuRepository

MONTHS = {
    "янв": 1,
    "фев": 2,
    "мар": 3,
    "апр": 4,
    "май": 5,
    "июн": 6,
    "июл": 7,
    "авг": 8,
    "сен": 9,
    "окт": 10,
    "ноя": 11,
    "дек": 12,
}


@dataclass(frozen=True, slots=True)
class LoaderConfig:
    """Порог выброса консервативен, чтобы не срезать обычный опт."""

    invoice_factor: float = 10.0
    invoice_share: float = 0.50


@dataclass(frozen=True, slots=True)
class LoadedDataset:
    repository: MemorySkuRepository
    report: DataQualityReport


@dataclass(frozen=True, slots=True)
class SupplierFiles:
    supplier: str
    monthly_sales: Path
    monthly_stock: Path
    transactions: Path
    moq: Path
    in_transit: Path
    seasonality: Path


@dataclass(slots=True)
class _Draft:
    code: str
    supplier: str
    name: str = ""
    article: str = ""
    category: str = ""
    unit: str = "шт"
    moq: int = 1
    on_hand_stock: float | None = None
    reserved_stock: float = 0.0
    free_stock: float = 0.0
    in_transit: float = 0.0
    inbound: list[Inbound] = field(default_factory=list)
    sales: dict[date, float] = field(default_factory=dict)
    stocks: dict[date, float] = field(default_factory=dict)
    stock_history_known: bool = False
    bulk_orders: tuple[BulkOrderEvent, ...] = ()


def load_dataset(data_dir: Path, config: LoaderConfig | None = None) -> LoadedDataset:
    """Прочитать обе выгрузки и вернуть доменные сущности с аудитом."""
    config = config or LoaderConfig()
    report = DataQualityReport()
    file_sets = discover_supplier_files(data_dir)
    report.files = [
        str(path.relative_to(data_dir)) for files in file_sets for path in _paths(files)
    ]

    drafts: dict[str, _Draft] = {}
    for files in file_sets:
        supplier_drafts = _load_supplier(files, report, config)
        for code, draft in supplier_drafts.items():
            if code in drafts and drafts[code].supplier != draft.supplier:
                raise StorageError(
                    f"Код 1С {code!r} встречается у двух поставщиков",
                    details={"suppliers": [drafts[code].supplier, draft.supplier]},
                )
            drafts[code] = draft

    skus = sorted((_to_sku(draft) for draft in drafts.values()), key=lambda x: (x.supplier, x.code))
    report.metrics.update(
        {
            "normalized_skus": len(skus),
            "suppliers": dict(Counter(sku.supplier for sku in skus)),
            "history_months_max": max((sku.months_of_history for sku in skus), default=0),
            "confirmed_bulk_orders": sum(len(sku.bulk_orders) for sku in skus),
            "skus_without_sales": sum(not sku.history for sku in skus),
        }
    )
    return LoadedDataset(MemorySkuRepository(skus), report)


def discover_supplier_files(data_dir: Path) -> list[SupplierFiles]:
    """Найти файлы по назначению, не полагаясь на лишний уровень каталогов."""
    if not data_dir.exists():
        raise StorageError(f"Каталог данных не найден: {data_dir}")

    result = []
    for supplier, marker in (("IEK", "iek"), ("Systeme Electric", "systeme")):
        paths = [path for path in data_dir.rglob("*.xlsx") if marker in str(path).casefold()]
        result.append(
            SupplierFiles(
                supplier=supplier,
                monthly_sales=_one(paths, "ежемесячные продажи", supplier),
                monthly_stock=_one(paths, "ежемесячные остатки", supplier),
                transactions=_one(paths, "динамика продаж", supplier),
                moq=_one(paths, "moq", supplier),
                in_transit=_one(
                    paths,
                    "путь иэк" if supplier == "IEK" else "товар в пути",
                    supplier,
                ),
                seasonality=_one(paths, "сезонность", supplier),
            )
        )
    return result


def _one(paths: list[Path], marker: str, supplier: str) -> Path:
    matches = [path for path in paths if marker in path.name.casefold()]
    if len(matches) != 1:
        raise StorageError(
            f"Для {supplier} ожидался один файл {marker!r}, найдено {len(matches)}",
            details={"matches": [str(path) for path in matches]},
        )
    return matches[0]


def _paths(files: SupplierFiles) -> tuple[Path, ...]:
    return (
        files.monthly_sales,
        files.monthly_stock,
        files.transactions,
        files.moq,
        files.in_transit,
        files.seasonality,
    )


def _load_supplier(
    files: SupplierFiles, report: DataQualityReport, config: LoaderConfig
) -> dict[str, _Draft]:
    drafts: dict[str, _Draft] = {}
    article_to_code: dict[str, str] = {}
    _read_monthly_sales(files, drafts, article_to_code, report)
    _read_monthly_stock(files, drafts, report)
    _read_moq(files, drafts, article_to_code, report)
    if files.supplier == "IEK":
        _read_iek_transit(files, drafts, article_to_code, report)
    else:
        _read_systeme_model(files, drafts, article_to_code, report)
    _read_transactions(files, drafts, report, config)
    _audit_completeness(files, drafts, report)
    return drafts


def _read_monthly_sales(
    files: SupplierFiles,
    drafts: dict[str, _Draft],
    article_to_code: dict[str, str],
    report: DataQualityReport,
) -> None:
    workbook = load_workbook(files.monthly_sales, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [_text(value) for value in next(rows)]
    next(rows, None)
    code_idx = _column(header, "Номенклатура.Код")
    name_idx = _column(header, "Номенклатура")
    article_idx = header.index("Артикул") if "Артикул" in header else None
    month_columns = _month_columns(header)
    seen: Counter[str] = Counter()
    negative_values = 0

    for row in rows:
        code = _code(_at(row, code_idx))
        if not code:
            continue
        seen[code] += 1
        draft = drafts.setdefault(code, _Draft(code=code, supplier=files.supplier))
        draft.name = draft.name or _text(_at(row, name_idx))
        if article_idx is not None:
            _set_article(draft, _text(_at(row, article_idx)), article_to_code, report)
        for idx, month in month_columns.items():
            value = _number(_at(row, idx), default=0.0)
            negative_values += value < 0
            draft.sales[month] = max(0.0, value)
    workbook.close()

    report.add_issue(
        "duplicate_monthly_sales_code",
        "error",
        "Код 1С повторяется в помесячных продажах",
        count=_duplicates(seen),
        supplier=files.supplier,
    )
    report.add_issue(
        "negative_monthly_sales",
        "warning",
        "Отрицательный месячный итог трактуется как возврат, а не спрос",
        count=negative_values,
        supplier=files.supplier,
    )
    report.metrics[f"{files.supplier}.sales_codes"] = len(seen)
    report.metrics[f"{files.supplier}.sales_months"] = len(month_columns)


def _read_monthly_stock(
    files: SupplierFiles, drafts: dict[str, _Draft], report: DataQualityReport
) -> None:
    workbook = load_workbook(files.monthly_stock, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [_text(value) for value in next(rows)]
    next(rows, None)
    next(rows, None)
    code_idx = _column(header, "Номенклатура.Код")
    name_idx = _column(header, "Номенклатура")
    unit_idx = _first_column(header, "Ед.", "Ед.изм")
    month_columns = _month_columns(header)
    seen: Counter[str] = Counter()
    blanks = negatives = 0

    for row in rows:
        code = _code(_at(row, code_idx))
        if not code:
            continue
        seen[code] += 1
        draft = drafts.setdefault(code, _Draft(code=code, supplier=files.supplier))
        draft.name = draft.name or _text(_at(row, name_idx))
        draft.unit = _text(_at(row, unit_idx)) or draft.unit
        draft.stock_history_known = True
        for idx, month in month_columns.items():
            raw = _at(row, idx)
            blanks += raw in (None, "")
            value = _number(raw, default=0.0)
            negatives += value < 0
            draft.stocks[month] = max(0.0, value)
        latest = max(month_columns.values())
        draft.on_hand_stock = draft.stocks[latest]
        draft.free_stock = draft.on_hand_stock
    workbook.close()

    report.add_issue(
        "duplicate_stock_code",
        "error",
        "Код 1С повторяется в истории остатков",
        count=_duplicates(seen),
        supplier=files.supplier,
    )
    report.add_issue(
        "blank_stock_month",
        "info",
        "Пустой начальный остаток трактуется как подтверждённый stockout",
        count=blanks,
        supplier=files.supplier,
    )
    report.add_issue(
        "negative_stock",
        "warning",
        "Отрицательный остаток нормализован до нуля",
        count=negatives,
        supplier=files.supplier,
    )
    report.metrics[f"{files.supplier}.stock_codes"] = len(seen)


def _read_moq(
    files: SupplierFiles,
    drafts: dict[str, _Draft],
    article_to_code: dict[str, str],
    report: DataQualityReport,
) -> None:
    workbook = load_workbook(files.moq, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [_text(value) for value in next(rows)]
    code_idx = _first_column(header, "Код 1с", "Номенклатура.Код")
    name_idx = _first_column(header, "Наименование", "Номенклатура")
    article_idx = _first_column(header, "Артикул поставщика", "Артикул")
    moq_idx = _first_column(header, "Мин. разр. к отгр.", "Кратность")
    seen: Counter[str] = Counter()
    invalid = 0

    for row in rows:
        code = _code(_at(row, code_idx))
        if not code:
            continue
        seen[code] += 1
        draft = drafts.setdefault(code, _Draft(code=code, supplier=files.supplier))
        draft.name = draft.name or _text(_at(row, name_idx))
        _set_article(draft, _text(_at(row, article_idx)), article_to_code, report)
        raw = _at(row, moq_idx)
        parsed = _positive_int(raw)
        invalid += parsed is None
        draft.moq = parsed or 1
    workbook.close()

    report.add_issue(
        "duplicate_moq_code",
        "warning",
        "Повторяющиеся строки MOQ сведены по коду 1С",
        count=_duplicates(seen),
        supplier=files.supplier,
    )
    report.add_issue(
        "invalid_moq",
        "warning",
        "Некорректная или нулевая кратность заменена на 1",
        count=invalid,
        supplier=files.supplier,
    )


def _read_iek_transit(
    files: SupplierFiles,
    drafts: dict[str, _Draft],
    article_to_code: dict[str, str],
    report: DataQualityReport,
) -> None:
    workbook = load_workbook(files.in_transit, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [_text(value) for value in next(rows)]
    code_idx = _column(header, "Код 1с")
    article_idx = _column(header, "Артикул ИЭК")
    name_idx = _first_column(header, "Наименование", " Наименование")
    inbound_columns = {
        idx: _parse_inbound_header(value) for idx, value in enumerate(header[3:], start=3)
    }
    seen: Counter[str] = Counter()
    mapped_by_article = 0

    for row in rows:
        code = _code(_at(row, code_idx))
        article = _text(_at(row, article_idx))
        if not code and article:
            code = article_to_code.get(_article_key(article), "")
            mapped_by_article += bool(code)
        if not code:
            continue
        seen[code] += 1
        draft = drafts.setdefault(code, _Draft(code=code, supplier=files.supplier))
        draft.name = draft.name or _text(_at(row, name_idx))
        _set_article(draft, article, article_to_code, report)
        for idx, (eta, document) in inbound_columns.items():
            quantity = max(0.0, _number(_at(row, idx), default=0.0))
            if quantity <= 0:
                continue
            draft.in_transit += quantity
            draft.inbound.append(Inbound(quantity=quantity, eta=eta, document=document))
    workbook.close()

    report.metrics[f"{files.supplier}.transit_mapped_by_article"] = mapped_by_article
    report.add_issue(
        "duplicate_transit_code",
        "info",
        "Несколько строк пути по коду суммированы",
        count=_duplicates(seen),
        supplier=files.supplier,
    )
    report.add_issue(
        "reserve_unavailable",
        "warning",
        "В IEK нет резерва; свободный остаток принят равным последнему остатку",
        supplier=files.supplier,
    )


def _read_systeme_model(
    files: SupplierFiles,
    drafts: dict[str, _Draft],
    article_to_code: dict[str, str],
    report: DataQualityReport,
) -> None:
    workbook = load_workbook(files.in_transit, read_only=True, data_only=True)
    rows = workbook["TDSheet"].iter_rows(values_only=True)
    next(rows, None)
    header = [_text(value) for value in next(rows)]
    code_idx = _column(header, "Код 1с")
    article_idx = _column(header, "Артикул поставщика")
    name_idx = _column(header, "Наименование")
    category_idx = _column(header, "Категория 2026")
    on_hand_idx = _column(header, "Остаток")
    reserve_idx = _column(header, "Зарезервировано")
    free_idx = _column(header, "Свободный остаток")
    transit_idx = next(
        idx for idx, value in enumerate(header) if value.casefold().startswith("сэ в пути")
    )
    inconsistent_free = 0

    for row in rows:
        code = _code(_at(row, code_idx))
        article = _text(_at(row, article_idx))
        if not code and article:
            code = article_to_code.get(_article_key(article), "")
        if not code:
            continue
        draft = drafts.setdefault(code, _Draft(code=code, supplier=files.supplier))
        draft.name = draft.name or _text(_at(row, name_idx))
        draft.category = _text(_at(row, category_idx))
        _set_article(draft, article, article_to_code, report)
        draft.on_hand_stock = max(0.0, _number(_at(row, on_hand_idx), default=0.0))
        draft.reserved_stock = max(0.0, _number(_at(row, reserve_idx), default=0.0))
        calculated_free = max(0.0, draft.on_hand_stock - draft.reserved_stock)
        draft.free_stock = max(0.0, _number(_at(row, free_idx), default=calculated_free))
        inconsistent_free += abs(draft.free_stock - calculated_free) > 0.01
        draft.in_transit = max(0.0, _number(_at(row, transit_idx), default=0.0))
        if draft.in_transit > 0:
            draft.inbound = [
                Inbound(quantity=draft.in_transit, eta=None, document="СЭ в пути")
            ]
    workbook.close()

    report.add_issue(
        "free_stock_mismatch",
        "warning",
        "Свободный остаток не равен остатку минус резерв; сохранено значение менеджера",
        count=inconsistent_free,
        supplier=files.supplier,
    )


def _read_transactions(
    files: SupplierFiles,
    drafts: dict[str, _Draft],
    report: DataQualityReport,
    config: LoaderConfig,
) -> None:
    workbook = load_workbook(files.transactions, read_only=True, data_only=True)
    rows = workbook.worksheets[0].iter_rows(values_only=True)
    header = [_text(value) for value in next(rows)]
    date_idx = _column(header, "Дата")
    invoice_idx = _column(header, "Номер")
    code_idx = _column(header, "Код")
    unit_idx = _column(header, "Ед.")
    quantity_idx = _column(header, "Количество")
    invoice_totals: dict[str, dict[tuple[str, date], float]] = defaultdict(
        lambda: defaultdict(float)
    )
    units_by_code: dict[str, Counter[str]] = defaultdict(Counter)
    seen_rows: set[tuple[str, str, str, float]] = set()
    positive = returns = invalid_dates = duplicate_rows = 0

    for row in rows:
        code = _code(_at(row, code_idx))
        quantity = _number(_at(row, quantity_idx), default=0.0)
        if not code or quantity == 0:
            continue
        date_text = _text(_at(row, date_idx))
        invoice = _text(_at(row, invoice_idx)) or "без номера"
        row_key = (date_text, invoice, code, quantity)
        duplicate_rows += row_key in seen_rows
        seen_rows.add(row_key)
        unit = _text(_at(row, unit_idx))
        if unit:
            units_by_code[code][unit] += 1
        if quantity < 0:
            returns += 1
            continue
        positive += 1
        month = _transaction_month(_at(row, date_idx))
        if month is None:
            invalid_dates += 1
            continue
        invoice_totals[code][(invoice, month)] += quantity
    workbook.close()

    bulk_count = 0
    for code, totals in invoice_totals.items():
        quantities = list(totals.values())
        if len(quantities) < 4:
            continue
        regular = statistics.median(quantities)
        total = sum(quantities)
        events = [
            BulkOrderEvent(month, invoice, quantity, regular)
            for (invoice, month), quantity in totals.items()
            if regular > 0
            and quantity >= regular * config.invoice_factor
            and quantity / total >= config.invoice_share
        ]
        if events:
            draft = drafts.setdefault(code, _Draft(code=code, supplier=files.supplier))
            draft.bulk_orders = tuple(
                sorted(events, key=lambda event: (event.month, event.invoice))
            )
            bulk_count += len(events)

    unit_mismatches = 0
    for code, units in units_by_code.items():
        draft = drafts.get(code)
        if draft is None:
            continue
        dominant = units.most_common(1)[0][0]
        if draft.unit == "шт" and dominant != "шт":
            draft.unit = dominant
        unit_mismatches += sum(count for unit, count in units.items() if unit != draft.unit)

    report.metrics[f"{files.supplier}.transaction_sales"] = positive
    report.metrics[f"{files.supplier}.transaction_returns"] = returns
    report.metrics[f"{files.supplier}.bulk_orders"] = bulk_count
    report.add_issue(
        "returns",
        "info",
        "Отрицательные транзакции исключены из поиска крупных заказов как возвраты",
        count=returns,
        supplier=files.supplier,
    )
    report.add_issue(
        "duplicate_transaction",
        "warning",
        "Полностью повторяющиеся транзакции требуют проверки",
        count=duplicate_rows,
        supplier=files.supplier,
    )
    report.add_issue(
        "unit_mismatch",
        "warning",
        "Единица транзакции отличается от основной единицы позиции",
        count=unit_mismatches,
        supplier=files.supplier,
    )
    report.add_issue(
        "invalid_transaction_date",
        "warning",
        "Неизвестная дата исключила транзакцию из клиентского сценария",
        count=invalid_dates,
        supplier=files.supplier,
    )


def _audit_completeness(
    files: SupplierFiles, drafts: dict[str, _Draft], report: DataQualityReport
) -> None:
    checks = (
        ("missing_sales_history", "Нет строки помесячных продаж", lambda item: not item.sales),
        (
            "missing_stock_history",
            "Нет строки истории остатка; stockout не подтверждается",
            lambda item: not item.stock_history_known,
        ),
        ("missing_supplier_article", "Не найден артикул поставщика", lambda item: not item.article),
    )
    for code, message, predicate in checks:
        report.add_issue(
            code,
            "warning",
            message,
            count=sum(predicate(draft) for draft in drafts.values()),
            supplier=files.supplier,
        )


def _to_sku(draft: _Draft) -> Sku:
    points = tuple(
        MonthPoint(
            month=month,
            sold=sold,
            stock_start=draft.stocks.get(month, 0.0),
            stock_known=draft.stock_history_known,
        )
        for month, sold in sorted(draft.sales.items())
    )
    return Sku(
        code=draft.code,
        name=draft.name or draft.code,
        supplier=draft.supplier,
        article=draft.article,
        category=draft.category,
        unit=draft.unit,
        moq=max(1, draft.moq),
        on_hand_stock=draft.on_hand_stock,
        reserved_stock=draft.reserved_stock,
        free_stock=max(0.0, draft.free_stock),
        in_transit=max(0.0, draft.in_transit),
        inbound=tuple(draft.inbound),
        history=points,
        bulk_orders=draft.bulk_orders,
    )


def _set_article(
    draft: _Draft,
    article: str,
    article_to_code: dict[str, str],
    report: DataQualityReport,
) -> None:
    if not article:
        return
    key = _article_key(article)
    known_code = article_to_code.get(key)
    if known_code and known_code != draft.code:
        report.add_issue(
            "article_code_conflict",
            "warning",
            f"Артикул {article!r} связан с несколькими кодами 1С",
            supplier=draft.supplier,
        )
        return
    article_to_code[key] = draft.code
    if draft.article and _article_key(draft.article) != key:
        report.add_issue(
            "code_article_conflict",
            "warning",
            f"Код 1С {draft.code!r} связан с разными артикулами",
            supplier=draft.supplier,
        )
        return
    draft.article = article


def _month_columns(header: list[str]) -> dict[int, date]:
    result = {
        idx: month
        for idx, value in enumerate(header)
        if (month := _parse_month(value)) is not None
    }
    if not result:
        raise StorageError("В Excel не найдены помесячные колонки")
    return result


def _parse_month(value: str) -> date | None:
    normalized = value.casefold().replace("ё", "е").replace("\xa0", " ")
    year_match = re.search(r"20\d{2}", normalized)
    if not year_match:
        return None
    for prefix, month in MONTHS.items():
        if prefix in normalized:
            return date(int(year_match.group()), month, 1)
    return None


def _transaction_month(value: Any) -> date | None:
    if isinstance(value, datetime):
        return date(value.year, value.month, 1)
    if isinstance(value, date):
        return value.replace(day=1)
    text = _text(value)
    for pattern in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            parsed = datetime.strptime(text, pattern)
            return date(parsed.year, parsed.month, 1)
        except ValueError:
            continue
    return None


def _parse_inbound_header(value: str) -> tuple[date | None, str]:
    """Извлечь ETA и узнаваемый номер документа из заголовка заказа ИЭК."""
    text = _text(value).replace("\\xa0", " ")
    eta_match = re.search(r"поступление до\s+(\d{2}\.\d{2}\.\d{4})", text, re.IGNORECASE)
    document_match = re.search(r"УТ-\d+", text, re.IGNORECASE)
    eta = datetime.strptime(eta_match.group(1), "%d.%m.%Y").date() if eta_match else None
    document = document_match.group(0).upper() if document_match else text
    return eta, document


def _column(header: list[str], name: str) -> int:
    if name not in header:
        raise StorageError(f"Обязательная колонка {name!r} не найдена")
    return header.index(name)


def _first_column(header: list[str], *names: str) -> int:
    for name in names:
        if name in header:
            return header.index(name)
    raise StorageError(f"Ни одна из обязательных колонок {names!r} не найдена")


def _at(row: tuple[Any, ...], idx: int) -> Any:
    return row[idx] if idx < len(row) else None


def _text(value: Any) -> str:
    return "" if value is None else str(value).replace("\xa0", " ").strip()


def _code(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _text(value)


def _article_key(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def _number(value: Any, *, default: float) -> float:
    if isinstance(value, bool):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(" ", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return default


def _positive_int(value: Any) -> int | None:
    number = _number(value, default=0.0)
    return max(1, round(number)) if number > 0 else None


def _duplicates(values: Counter[str]) -> int:
    return sum(count - 1 for count in values.values() if count > 1)
