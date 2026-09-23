"""Таблица заказа в CSV или xlsx с произвольным набором колонок.

Состав колонок выбирает менеджер: образца файла импорта 1С у нас нет, и
подогнать выгрузку под него должно быть можно без правки кода.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

# Excel запрещает эти символы в имени листа и режет его до 31 знака
FORBIDDEN = str.maketrans({c: " " for c in "[]:*?/\\"})


def render_csv(header: list[str], rows: list[list[Any]], *, separator: str,
               encoding: str) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=separator)
    writer.writerow(header)
    writer.writerows(rows)
    # cp1251 не знает части символов (например, «№» знает, а эмодзи нет) — заменяем
    return buffer.getvalue().encode(encoding, errors="replace")


def render_xlsx(header: list[str], sheets: dict[str, list[list[Any]]]) -> bytes:
    """Лист на поставщика: заказ поставщику в 1С — один документ на контрагента."""
    wb = Workbook()
    wb.remove(wb.active)
    # Пустой заказ — всё равно валидный файл с шапкой, а не битый xlsx без листов
    for name, rows in (sheets or {"Заказ": []}).items():
        ws = wb.create_sheet(name.translate(FORBIDDEN)[:31])
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
