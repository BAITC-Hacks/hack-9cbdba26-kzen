"""Сравнение формулы менеджера и нашего расчёта по всему ассортименту.

Это не оценка точности прогноза: будущего факта для честного backtest пока нет.
Скрипт измеряет воспроизводимый контрфактический эффект очистки подтверждённых
крупных накладных и компенсации stockout при одинаковых остатках и параметрах.
"""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from dataclasses import dataclass, replace
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from app.application.use_cases.replenishment import (  # noqa: E402
    CATEGORY_SAFETY_UPLIFT,
    CalcParams,
    _round_to_moq,
    _split_inbound,
    calculate_order_line,
)
from app.domain.demand import (  # noqa: E402
    compensate_stockouts,
    remove_bulk_orders,
    remove_confirmed_bulk_orders,
)
from app.domain.entities import MonthPoint, Sku  # noqa: E402
from app.infrastructure.forecasting.baseline import BaselineForecaster  # noqa: E402
from app.infrastructure.forecasting.smoothed import SmoothedForecaster  # noqa: E402
from app.infrastructure.storage.excel_loader import load_dataset  # noqa: E402


@dataclass
class Effect:
    positions: int = 0
    nominal_quantity: int = 0
    cost_tenge: float = 0.0
    priced_positions: int = 0

    def add(self, sku: Sku, quantity: int, costs: dict[str, float]) -> None:
        self.positions += 1
        self.nominal_quantity += quantity
        if sku.supplier == "Systeme Electric" and sku.code in costs:
            self.cost_tenge += quantity * costs[sku.code]
            self.priced_positions += 1


def _code(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return "" if value is None else str(value).strip()


def _systeme_costs(data_dir: Path) -> dict[str, float]:
    candidates = list(data_dir.rglob("*Товар в пути*SystemElectric*.xlsx"))
    if len(candidates) != 1:
        raise RuntimeError(f"Ожидался один TDSheet Systeme, найдено {len(candidates)}")

    workbook = load_workbook(candidates[0], read_only=True, data_only=True)
    rows = workbook["TDSheet"].iter_rows(values_only=True)
    next(rows, None)
    header = ["" if value is None else str(value).strip() for value in next(rows)]
    code_idx = header.index("Код 1с")
    cost_idx = header.index("СС реал")
    result = {}
    for row in rows:
        code = _code(row[code_idx])
        cost = row[cost_idx]
        if code and isinstance(cost, (int, float)) and cost >= 0:
            result[code] = float(cost)
    workbook.close()
    return result


def _params(sku: Sku, today: date) -> CalcParams:
    lead_time = 30 if sku.supplier == "Systeme Electric" else 35
    return CalcParams(
        lead_time_days=lead_time,
        coverage_months=14 / 30,
        safety_factor=0.2,
        today=today,
    )


def _quantity_for_history(
    sku: Sku,
    history: tuple[MonthPoint, ...],
    params: CalcParams,
    forecaster: BaselineForecaster,
) -> int:
    """Один и тот же заказ для разных стадий очистки истории."""
    lead_days = sku.lead_time_days or params.lead_time_days
    arrival = (params.today or date.today()) + timedelta(days=lead_days)
    forecast = forecaster.forecast(replace(sku, history=history, bulk_orders=()), arrival.month)
    lead_months = lead_days / 30
    horizon = lead_months + params.coverage_months
    category_uplift = (
        CATEGORY_SAFETY_UPLIFT.get(sku.category, 0.0)
        if sku.supplier == "Systeme Electric"
        else 0.0
    )
    safety = min(1.0, params.safety_factor + category_uplift)
    need = forecast.monthly_demand * horizon * (1 + safety)
    counted_inbound, _ = _split_inbound(sku, arrival)
    return _round_to_moq(need - sku.free_stock - counted_inbound, sku.moq)


def _effect_dict(effect: Effect, by_unit: dict[str, int]) -> dict[str, Any]:
    return {
        "affected_positions": effect.positions,
        "nominal_quantity": effect.nominal_quantity,
        "quantity_by_unit": dict(sorted(by_unit.items())),
        "systeme_excess_cost_tenge": round(effect.cost_tenge, 2),
        "systeme_priced_positions": effect.priced_positions,
    }


def evaluate(data_dir: Path, today: date) -> tuple[dict[str, Any], dict[str, Any]]:
    loaded = load_dataset(data_dir)
    skus = loaded.repository.list_skus()
    costs = _systeme_costs(data_dir)
    baseline = BaselineForecaster()
    smoothed = SmoothedForecaster()

    invoice_effect = Effect()
    stockout_effect = Effect()
    all_bulk_effect = Effect()
    invoice_by_unit: dict[str, int] = defaultdict(int)
    stockout_by_unit: dict[str, int] = defaultdict(int)
    all_bulk_by_unit: dict[str, int] = defaultdict(int)
    baseline_total = smoothed_total = 0
    baseline_higher = smoothed_higher = equal = 0
    confirmed_positions = 0
    confirmed_events = 0

    for sku in skus:
        params = _params(sku, today)
        baseline_line = calculate_order_line(sku, params, baseline)
        smoothed_line = calculate_order_line(sku, params, smoothed)
        baseline_total += baseline_line.quantity
        smoothed_total += smoothed_line.quantity
        if baseline_line.quantity > smoothed_line.quantity:
            baseline_higher += 1
        elif smoothed_line.quantity > baseline_line.quantity:
            smoothed_higher += 1
        else:
            equal += 1

        confirmed, _, _, _ = remove_confirmed_bulk_orders(sku.history, sku.bulk_orders)
        without_bulk, _, _ = remove_bulk_orders(confirmed)
        compensated, _, _ = compensate_stockouts(without_bulk)

        raw_quantity = _quantity_for_history(sku, sku.history, params, baseline)
        after_invoices = _quantity_for_history(sku, confirmed, params, baseline)
        after_all_bulk = _quantity_for_history(sku, without_bulk, params, baseline)
        after_stockout = _quantity_for_history(sku, compensated, params, baseline)

        if sku.bulk_orders:
            confirmed_positions += 1
            confirmed_events += len(sku.bulk_orders)
        invoice_delta = max(0, raw_quantity - after_invoices)
        if invoice_delta:
            invoice_effect.add(sku, invoice_delta, costs)
            invoice_by_unit[sku.unit] += invoice_delta

        all_bulk_delta = max(0, raw_quantity - after_all_bulk)
        if all_bulk_delta:
            all_bulk_effect.add(sku, all_bulk_delta, costs)
            all_bulk_by_unit[sku.unit] += all_bulk_delta

        stockout_delta = max(0, after_stockout - after_all_bulk)
        if stockout_delta:
            stockout_effect.add(sku, stockout_delta, costs)
            stockout_by_unit[sku.unit] += stockout_delta

        if baseline_line.quantity != after_stockout:
            raise AssertionError(f"Контрфактический расчёт разошёлся с baseline: {sku.code}")

    metrics = {
        "task": "supplier_replenishment_business_effect",
        "scope": {
            "skus": len(skus),
            "suppliers": loaded.repository.suppliers(),
            "data_period": "2024-01..2026-09",
        },
        "primary_metric": "counterfactual_reorder_quantity",
        "direction": "reduce one-off overorder and stockout underorder",
        "comparison": {
            "baseline_total_nominal_quantity": baseline_total,
            "smoothed_total_nominal_quantity": smoothed_total,
            "baseline_higher_positions": baseline_higher,
            "smoothed_higher_positions": smoothed_higher,
            "equal_positions": equal,
        },
        "confirmed_invoice_candidates": {
            "candidate_positions": confirmed_positions,
            "candidate_events": confirmed_events,
            **_effect_dict(invoice_effect, invoice_by_unit),
        },
        "all_bulk_outliers": _effect_dict(all_bulk_effect, all_bulk_by_unit),
        "stockout_compensation": _effect_dict(stockout_effect, stockout_by_unit),
        "forecast_accuracy": None,
        "recommendation": "KEEP BASELINE",
        "reason": (
            "Правила выбросов и stockout дают измеримый контрфактический эффект; "
            "превосходство smoothed по точности требует хронологического backtest."
        ),
    }
    config = {
        "task": "deterministic inventory replenishment",
        "data_dir": (
            str(data_dir.relative_to(ROOT)) if data_dir.is_relative_to(ROOT) else str(data_dir)
        ),
        "data_files": sorted(str(path.relative_to(data_dir)) for path in data_dir.rglob("*.xlsx")),
        "calculation_date": today.isoformat(),
        "split": {
            "type": "none",
            "status": "not_run",
            "reason": "Контрфактический аудит полного ассортимента, не оценка точности прогноза.",
        },
        "compared_approaches": ["baseline", "smoothed"],
        "counterfactuals": [
            "raw vs confirmed-invoice-cleaned history",
            "bulk-cleaned vs stockout-compensated history",
        ],
        "defaults": {
            "lead_time_days": {"IEK": 35, "Systeme Electric": 30},
            "coverage_months": 14 / 30,
            "safety_factor": 0.2,
        },
        "money": {
            "supplier": "Systeme Electric",
            "source": "TDSheet.СС реал",
            "currency": "KZT",
        },
        "random_seed": None,
    }
    return metrics, config


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--data-dir", type=Path, default=ROOT / "data")
    parser.add_argument("--today", type=date.fromisoformat, default=date(2026, 9, 23))
    parser.add_argument("--metrics", type=Path, default=ROOT / "metrics.json")
    parser.add_argument("--config", type=Path, default=ROOT / "run_config.json")
    args = parser.parse_args()

    metrics, config = evaluate(args.data_dir, args.today)
    args.metrics.write_text(json.dumps(metrics, ensure_ascii=False, indent=2) + "\n")
    args.config.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n")

    invoice = metrics["confirmed_invoice_candidates"]
    stockout = metrics["stockout_compensation"]
    comparison = metrics["comparison"]
    print(f"SKU: {metrics['scope']['skus']}")
    print(
        "Разовые подтверждённые накладные: "
        f"{invoice['affected_positions']} поз., {invoice['nominal_quantity']:,} ед.; "
        f"Systeme: {invoice['systeme_excess_cost_tenge']:,.0f} ₸"
    )
    print(
        "Игнорирование stockout: "
        f"{stockout['affected_positions']} поз., "
        f"{stockout['nominal_quantity']:,} ед. недозаказа"
    )
    print(
        "baseline / smoothed: "
        f"{comparison['baseline_total_nominal_quantity']:,} / "
        f"{comparison['smoothed_total_nominal_quantity']:,} номинальных единиц"
    )
    print(f"Рекомендация: {metrics['recommendation']} — {metrics['reason']}")


if __name__ == "__main__":
    main()
