"""Агент закупщика: ветвления по данным, what-if с пересчётом, выгрузка для 1С."""

from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tests.conftest import history

from app.core.container import Container
from app.domain.entities import Sku
from app.infrastructure.cache.memory import MemoryCache
from app.infrastructure.forecasting.baseline import BaselineForecaster
from app.infrastructure.forecasting.smoothed import SmoothedForecaster
from app.infrastructure.storage.memory_repo import MemorySkuRepository
from app.main import create_app

AGENT = "/api/v1/agent/runs"
DRAFTS = "/api/v1/drafts"


@pytest.fixture
def agent_client(settings):
    spike = [40.0] * 24
    spike[20] = 4000.0  # одна гигантская продажа — клиентский сценарий Must have 4
    repo = MemorySkuRepository([
        # обычная позиция: нужна поставка, данные полные
        Sku(code="300200428_", name="Розетка A398", supplier="Systeme Electric",
            article="ATN000343", moq=10, free_stock=20.0, history=history([40.0] * 24)),
        # в пути достаточно, чтобы сейчас не заказывать; задержка это ломает
        Sku(code="300200430_", name="Выключатель A384", supplier="Systeme Electric",
            article="ATN000312", moq=1, free_stock=10.0, in_transit=200.0,
            history=history([40.0] * 24)),
        # артикул поставщика вместо кода 1С — как 6 позиций в модели менеджера
        Sku(code="ATN000330", name="Розетка USB", supplier="Systeme Electric",
            article="ATN000330", history=history([30.0] * 24)),
        # разовая крупная продажа в истории
        Sku(code="200400085_", name="F/UTP кат.5Е", supplier="IEK", article="LC1-C5E04-311",
            free_stock=0.0, history=history(spike)),
    ])
    container = Container(
        settings=settings, repo=repo, cache=MemoryCache(),
        forecasters={"baseline": BaselineForecaster(), "smoothed": SmoothedForecaster()},
    )
    with TestClient(create_app(settings, container=container)) as client:
        yield client


def _run(client, **payload) -> dict:
    response = client.post(AGENT, json={"goal": "Подготовь закупку", **payload})
    assert response.status_code == 201, response.text
    return response.json()


def _lines(draft: dict) -> dict[str, dict]:
    return {x["code"]: x for s in draft["suppliers"] for x in s["lines"]}


def _approve(client, draft: dict) -> dict:
    return client.post(f"{DRAFTS}/{draft['version']}/approve",
                       json={"author": "Алдияр", "revision": draft["revision"]}).json()


def test_agent_trace_shows_tools_and_waits_for_approval(agent_client):
    draft = _run(agent_client)
    tools = [e["tool"] for e in draft["trace"] if e["tool"]]

    assert tools[:3] == ["inspect_data_quality", "calculate_orders", "screen_recommendations"]
    assert draft["trace"][-1]["type"] == "awaiting_approval"
    assert draft["status"] == "draft"
    assert draft["sent_to_supplier"] is False


def test_missing_1c_code_is_insufficient_data_not_zero(agent_client):
    """Ветка B: не притворяемся, что всё знаем."""
    line = _lines(_run(agent_client))["ATN000330"]
    assert line["status"] == "insufficient_data"
    assert any("код 1С" in issue for issue in line["issues"])


def test_spike_triggers_extra_check_and_review(agent_client):
    """Ветка A: всплеск → агент сам добавляет разбор истории → позиция на проверку."""
    draft = _run(agent_client)
    reviews = [e for e in draft["trace"] if e["tool"] == "review_one_off"]

    assert [e["sku"] for e in reviews] == ["200400085_"]
    assert reviews[0]["facts"]["share_of_history"] > 0.3
    assert _lines(draft)["200400085_"]["status"] == "needs_review"


def test_inbound_delay_recalculates_and_requires_new_approval(agent_client):
    """Ветка C: поставка не успевает → пересчёт → новая версия → утверждать заново."""
    draft = _run(agent_client)
    assert "300200430_" not in _lines(draft)  # в пути хватает — заказывать не нужно
    approved = _approve(agent_client, draft)
    assert approved["status"] == "approved"

    response = agent_client.post(
        f"{AGENT}/{draft['version']}/inbound-delay",
        json={"code": "300200430_", "delayed_quantity": 200, "new_eta": "20.11.2026"},
    )
    assert response.status_code == 200
    derived = response.json()

    assert derived["version"] == draft["version"] + 1
    assert derived["parent_version"] == draft["version"]
    assert derived["status"] == "draft"
    assert _lines(derived)["300200430_"]["quantity"] > 0

    titles = [e["title"] for e in derived["trace"]]
    assert any("больше не актуально" in t for t in titles)
    recalculated = next(e for e in derived["trace"] if e["tool"] == "calculate_reorder_quantity")
    assert recalculated["facts"]["before"] == 0
    assert recalculated["facts"]["after"] > 0

    # старая утверждённая версия не тронута — видно, что утверждали
    assert agent_client.get(f"{DRAFTS}/{draft['version']}").json()["status"] == "approved"


def test_delay_that_changes_nothing_keeps_version(agent_client):
    draft = _run(agent_client)
    derived = agent_client.post(
        f"{AGENT}/{draft['version']}/inbound-delay",
        json={"code": "300200430_", "delayed_quantity": 1, "new_eta": "01.10.2026"},
    ).json()
    assert derived["version"] == draft["version"]


def test_xlsx_for_1c_has_sheet_per_supplier_and_skips_unknown(agent_client):
    draft = _run(agent_client)
    _approve(agent_client, draft)

    response = agent_client.get(f"{DRAFTS}/{draft['version']}/export")
    assert response.status_code == 200
    assert "purchase_order_1c" in response.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(response.content))
    assert "Сводка" in wb.sheetnames
    ws = wb["Systeme Electric"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Код 1с", "Артикул поставщика", "Номенклатура"]
    # выход по ТЗ п. 5: обоснование и срочность у каждой строки
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    explanation, urgency = header.index("Обоснование"), header.index("Срочность")
    assert all(row[explanation] and "заказать" in row[explanation] for row in rows)
    assert all(row[urgency] for row in rows)
    codes = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
    assert "300200428_" in codes
    assert "ATN000330" not in codes  # без кода 1С в загрузку не идёт


def test_unknown_sku_in_delay_is_404(agent_client):
    draft = _run(agent_client)
    response = agent_client.post(
        f"{AGENT}/{draft['version']}/inbound-delay",
        json={"code": "нет-такого", "delayed_quantity": 1, "new_eta": "01.10.2026"},
    )
    assert response.status_code == 404
