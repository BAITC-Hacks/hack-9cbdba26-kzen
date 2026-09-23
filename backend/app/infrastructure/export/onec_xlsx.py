"""Утверждённый заказ в xlsx для дальнейшей работы в 1С.

Колонки взяты из выгрузок партнёра, а не придуманы: «Код 1с» — ключ, по которому
1С сопоставит номенклатуру; «Артикул поставщика» — чтобы проверить глазами.
Лист на поставщика: заказ поставщику в 1С — это один документ на одного контрагента.

Сервис формирует файл, но никуда его не отправляет — это решает человек.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font

from app.domain.workflow import OrderDraft

COLUMNS = [
    "Код 1с", "Артикул поставщика", "Номенклатура", "Ед.", "Рекомендовано",
    "Количество", "Склад", "Причина корректировки", "Версия расчёта", "Утвердил",
]
WIDTHS = [14, 22, 48, 6, 14, 12, 12, 34, 14, 16]
# Excel запрещает эти символы в имени листа и режет его до 31 знака
FORBIDDEN = str.maketrans({c: " " for c in "[]:*?/\\"})


def render_purchase_order(draft: OrderDraft, *, warehouse: str,
                          supplier: str | None = None) -> bytes:
    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"

    exported = 0
    for name, lines in sorted(draft.by_supplier().items()):
        if supplier and name != supplier:
            continue
        rows = [x for x in lines if x.exportable]
        if not rows:
            continue
        ws = wb.create_sheet(name.translate(FORBIDDEN)[:31])
        ws.append(COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for line in rows:
            ws.append([
                line.code, line.article, line.name, line.unit, line.recommended,
                line.quantity, warehouse, line.adjustment_reason,
                f"v{draft.version}.r{draft.revision}", draft.approved_by,
            ])
        for i, width in enumerate(WIDTHS):
            ws.column_dimensions[chr(ord("A") + i)].width = width
        ws.freeze_panes = "A2"
        exported += len(rows)

    skipped = sum(1 for x in draft.lines.values() if x.quantity > 0 and not x.exportable)
    for row in [
        ("Версия расчёта", f"v{draft.version}.r{draft.revision}"),
        ("Утвердил", draft.approved_by),
        ("Дата утверждения", draft.approved_at.strftime("%d.%m.%Y %H:%M") if draft.approved_at
         else ""),
        ("Позиций в файле", exported),
        ("Не выгружено: недостаточно данных", skipped),
        ("Отправлено поставщику", "нет — заказ оформляет менеджер в 1С"),
    ]:
        summary.append(row)
    summary.column_dimensions["A"].width = 36
    summary.column_dimensions["B"].width = 40

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
